import os
import re
from openpyxl import load_workbook

def update_excel_with_screenshots(screenshots_dir, excel_path):
    """
    更新 Excel 文件中：
      - “截图文件名”列：填入文件名（按时间戳排序）
      - “测试用例编号”列：提取文件名中间部分
      - “需求编号”列：从测试用例编号提取 Rxxx
    ✅ 不改变原样式，仅更新单元格的 value。
    """

    # ===== 1️⃣ 检查文件与文件夹 =====
    if not os.path.isfile(excel_path):
        print(f"❌ 未找到 Excel 文件：{excel_path}")
        return
    if not os.path.isdir(screenshots_dir):
        print(f"❌ 未找到截图文件夹：{screenshots_dir}")
        return

    # ===== 2️⃣ 获取截图文件并排序 =====
    files = [
        f for f in os.listdir(screenshots_dir)
        if os.path.isfile(os.path.join(screenshots_dir, f))
        and f.lower().endswith(('.png', '.jpg', '.jpeg', '.bmp', '.gif'))
    ]

    def extract_timestamp(filename):
        match = re.match(r'^(\d+)', filename)
        return int(match.group(1)) if match else float('inf')

    files.sort(key=extract_timestamp)

    # ===== 3️⃣ 打开 Excel =====
    wb = load_workbook(excel_path)
    ws = wb.active

    # ===== 4️⃣ 找到目标列的列号 =====
    header_row = 1
    col_screenshot = None
    col_case_id = None
    col_req_id = None

    for col in range(1, ws.max_column + 1):
        val = str(ws.cell(row=header_row, column=col).value).strip() if ws.cell(row=header_row, column=col).value else ""
        if val == "截图文件名":
            col_screenshot = col
        elif val == "测试用例编号":
            col_case_id = col
        elif val == "需求编号":
            col_req_id = col

    if not col_screenshot:
        print("❌ 未找到列名 “截图文件名” ，无法写入。")
        return

    # ===== 5️⃣ 写入新内容（不清空样式，仅更新单元格值） =====
    for i, filename in enumerate(files, start=2):
        # 截图文件名
        ws.cell(row=i, column=col_screenshot).value = filename

        # 测试用例编号（去掉时间戳和扩展名）
        case_id = ""
        match = re.match(r'^\d+_(.*?)\.[^.]+$', filename)
        if match:
            case_id = match.group(1)

        if col_case_id:
            ws.cell(row=i, column=col_case_id).value = case_id

        # 需求编号（提取 R 开头部分）
        req_id = ""
        if case_id:
            match2 = re.search(r'(R\d+)', case_id)
            if match2:
                req_id = match2.group(1)

        if col_req_id:
            ws.cell(row=i, column=col_req_id).value = req_id

    # ===== 6️⃣ 保存文件 =====
    wb.save(excel_path)

    print(f"✅ 已成功更新：{excel_path}")
    print(f"📸 共写入 {len(files)} 个截图文件。")
    print("✅ 原表格样式完全保留。")


# =============================
# 使用示例
# =============================
if __name__ == "__main__":
    screenshots_dir = r"D:\000Desktop\VSCode\TestCtripFlight\screenshots"           # ← 你的截图文件夹路径
    excel_path = r"D:\000Desktop\Download\1763184687568_2025全国大学生软件测试大赛Web应用测试省赛\2025全国大学生软件测试大赛Web应用测试省赛\测试用例模版.xlsx"                      # ← 你的Excel文件路径

    update_excel_with_screenshots(screenshots_dir, excel_path)

